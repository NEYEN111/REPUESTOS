"""Laboratorio aislado: importaciones, artefactos inmutables y evaluación temporal."""
import csv
import hashlib
import io
import json
import math
import os
import re
import sqlite3
import subprocess
import threading
import uuid
import zipfile
from contextlib import contextmanager
from datetime import date, timedelta
from pathlib import Path
from backend.db import ROOT, now
from backend.services.predictions import runtime

FOLDER = Path(os.getenv('ML_LAB_DIR', str(ROOT/'data/ml/laboratorio')))
EXAMPLES = ROOT/'ml/examples/carparts'
SIMULATION = ROOT/'ml/examples/roda-ensayo'
LOCK = threading.Lock()


@contextmanager
def connection():
    FOLDER.mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(FOLDER/'laboratorio.db', timeout=30)
    c.row_factory = sqlite3.Row
    c.execute('PRAGMA foreign_keys=ON')
    try:
        yield c
        c.commit()
    except Exception:
        c.rollback()
        raise
    finally:
        c.close()


def init():
    with connection() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS datasets(id TEXT PRIMARY KEY,name TEXT,frequency TEXT,source TEXT,
          demo INTEGER DEFAULT 0,archived INTEGER DEFAULT 0,active_version TEXT,created_at TEXT);
        CREATE TABLE IF NOT EXISTS batches(id TEXT PRIMARY KEY,dataset TEXT REFERENCES datasets(id),
          filename TEXT,sha256 TEXT,records TEXT,summary TEXT,status TEXT,created_at TEXT,user_id INTEGER);
        CREATE TABLE IF NOT EXISTS observations(dataset TEXT REFERENCES datasets(id),sku TEXT,period TEXT,
          units INTEGER,batch TEXT REFERENCES batches(id),PRIMARY KEY(dataset,sku,period));
        CREATE TABLE IF NOT EXISTS versions(id TEXT PRIMARY KEY,dataset TEXT REFERENCES datasets(id),
          label TEXT,bundle TEXT,report TEXT,data_hash TEXT,created_at TEXT);
        CREATE TABLE IF NOT EXISTS jobs(id TEXT PRIMARY KEY,dataset TEXT REFERENCES datasets(id),action TEXT,
          status TEXT,result TEXT,error TEXT,created_at TEXT,finished_at TEXT,user_id INTEGER);
        CREATE TABLE IF NOT EXISTS activity(id INTEGER PRIMARY KEY,dataset TEXT,action TEXT,detail TEXT,
          created_at TEXT,user_id INTEGER);
        CREATE TABLE IF NOT EXISTS scenarios(dataset TEXT PRIMARY KEY REFERENCES datasets(id),
          kind TEXT NOT NULL, provenance_hash TEXT NOT NULL);
        """)
        # Older demonstrations activated V1 automatically even when a simple
        # historical average was more accurate. Keep the version, revoke that label.
        for item in c.execute('SELECT id,active_version FROM datasets WHERE demo=1 AND active_version IS NOT NULL').fetchall():
            version = row(c, 'SELECT report FROM versions WHERE id=? AND dataset=?',
                          (item['active_version'], item['id']))
            if not version:
                continue
            report = json.loads(version['report'])
            scores = report.get('candidates') or []
            baseline = (report.get('validation_baseline') or {}).get('mae')
            if baseline is not None and (not scores or min(s['mae'] for s in scores) >= baseline * .99):
                c.execute('UPDATE datasets SET active_version=NULL WHERE id=?', (item['id'],))
                event(c, item['id'], 'retirar_activacion_legacy',
                      {'version': item['active_version'], 'reason': 'No superó la referencia histórica en validación.'}, None)


def recover_jobs():
    init()
    with connection() as c:
        c.execute("UPDATE jobs SET status='error',error='El servidor se reinició. Inicia nuevamente el cálculo.' WHERE status='running'")


def row(c, sql, args=()):
    value = c.execute(sql, args).fetchone()
    return dict(value) if value else None


def dataset(c, did):
    d = row(c, 'SELECT * FROM datasets WHERE id=?', (did,))
    if not d:
        raise ValueError('Historial no encontrado.')
    return d


def event(c, did, action, detail, uid):
    c.execute('INSERT INTO activity(dataset,action,detail,created_at,user_id) VALUES(?,?,?,?,?)',
              (did, action, json.dumps(detail, ensure_ascii=False), now(), uid))


def create(name, frequency, source, uid, demo=False):
    if frequency not in ['monthly', 'daily'] or not name.strip() or not source.strip():
        raise ValueError('Indica nombre, frecuencia y procedencia del historial.')
    did = uuid.uuid4().hex
    with connection() as c:
        c.execute('INSERT INTO datasets VALUES(?,?,?,?,?,0,NULL,?)',
                  (did, name[:120], frequency, source[:1000], int(demo), now()))
        event(c, did, 'crear_historial', {'source': source[:1000]}, uid)
    return did


def source():
    path = EXAMPLES/'provenance.json'
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else None


def read_file(content, filename):
    if len(content) > 8*1024*1024:
        raise ValueError('El archivo supera 8 MB.')
    suffix = Path(filename).suffix.lower()
    if suffix == '.csv':
        try:
            text = content.decode('utf-8-sig')
            dialect = csv.Sniffer().sniff(text[:5000], delimiters=',;\t')
            values = list(csv.reader(io.StringIO(text), dialect))
        except (UnicodeError, csv.Error):
            raise ValueError('Guarda el CSV con codificación UTF-8 y columnas separadas por coma o punto y coma.')
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist()) > 40*1024*1024:
                    raise ValueError('El Excel descomprimido supera 40 MB.')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            if len(book.sheetnames) != 1:
                raise ValueError('El Excel debe contener una sola hoja de datos.')
            values = []
            for cells in book.active.iter_rows():
                if len(cells) > 30:
                    raise ValueError('El Excel debe contener como máximo 30 columnas.')
                if any(cell.data_type == 'f' for cell in cells):
                    raise ValueError('Usa valores guardados, no fórmulas, en el archivo de importación.')
                values.append([cell.value for cell in cells])
                if len(values) > 150001:
                    raise ValueError('El archivo supera 150,000 filas.')
            book.close()
        except (zipfile.BadZipFile, KeyError):
            raise ValueError('El archivo Excel no es válido.')
    else:
        raise ValueError('Selecciona un archivo .csv o .xlsx.')
    if len(values) < 2 or len(values) > 150001:
        raise ValueError('Se requieren entre 1 y 150,000 filas de datos.')
    headers = [str(x or '').strip().lower() for x in values[0]]
    if len(headers) > 30 or len(set(headers)) != len(headers):
        raise ValueError('Revisa columnas duplicadas o más de 30 columnas.')
    return headers, values[1:]


def preview(did, content, filename, mapping, uid):
    if did == 'roda-operacion-diaria':
        raise ValueError('Este historial recibe los cierres de la tienda automáticamente. Importa Excel en otro historial.')
    with connection() as c:
        d = dataset(c, did)
        if d['archived']:
            raise ValueError('Este historial está archivado; selecciona uno activo.')
        existing = {(r['sku'], r['period']): r['units'] for r in c.execute('SELECT * FROM observations WHERE dataset=?', (did,))}
    headers, values = read_file(content, filename)
    mapping = mapping or {'periodo': 'periodo', 'sku': 'sku', 'unidades': 'unidades'}
    if any(mapping.get(k) not in headers for k in ['periodo', 'sku', 'unidades']):
        return dict(headers=headers, needs_mapping=True, preview=values[:5])
    if len({mapping[k] for k in ['periodo', 'sku', 'unidades']}) != 3:
        raise ValueError('Cada dato debe corresponder a una columna diferente.')
    idx = {k: headers.index(mapping[k]) for k in ['periodo', 'sku', 'unidades']}
    errors, records, seen = [], [], {}
    duplicates = missing = conflicts = invalid = 0
    for line, values_row in enumerate(values, 2):
        if not any(v is not None and str(v).strip() for v in values_row):
            continue
        try:
            raw = {k: values_row[i] if i < len(values_row) else None for k, i in idx.items()}
            sku = str(raw['sku'] or '').strip()
            if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_. /-]{0,79}', sku):
                raise ValueError('SKU vacío o con caracteres no permitidos.')
            value = raw['periodo']
            if hasattr(value, 'strftime'):
                value = value.strftime('%Y-%m' if d['frequency'] == 'monthly' else '%Y-%m-%d')
            period = str(value or '').strip()
            if not re.fullmatch(r'\d{4}-\d{2}' if d['frequency'] == 'monthly' else r'\d{4}-\d{2}-\d{2}', period):
                raise ValueError('Periodo inválido: usa AAAA-MM para mensual o AAAA-MM-DD para diario.')
            parsed = date.fromisoformat(period+'-01' if len(period) == 7 else period)
            if parsed > date.fromisoformat(now()[:10]):
                raise ValueError('Una observación no puede tener fecha futura.')
            if period == now()[:7 if d['frequency'] == 'monthly' else 10]:
                raise ValueError('El periodo todavía no terminó. Importa solo meses o días completos anteriores al actual.')
            u = raw['unidades']
            units = None if u is None or str(u).strip() in ['', '?'] else float(u)
            if units is not None and (not math.isfinite(units) or units < 0 or units != int(units) or units > 1000000):
                raise ValueError('Unidades debe ser un entero entre 0 y 1,000,000 o vacío (desconocido).')
            units = int(units) if units is not None else None
            key = (sku, period)
            if key in seen or key in existing:
                previous = seen[key] if key in seen else existing[key]
                if previous != units:
                    conflicts += 1
                    raise ValueError('El producto/periodo ya existe con otra cantidad; no se sobrescribe.')
                duplicates += 1
                continue
            seen[key] = units
            records.append(dict(sku=sku, period=period, units=units))
            missing += units is None
        except (ValueError, TypeError) as e:
            invalid += 1
            if len(errors) < 30:
                errors.append(dict(row=line, message=str(e)))
    summary = dict(rows=len(values), accepted=len(records), duplicates=duplicates, missing=missing,
                   invalid=invalid, conflicts=conflicts, errors=errors, headers=headers,
                   products=len({r['sku'] for r in records}),
                   start=min((r['period'] for r in records), default=None),
                   end=max((r['period'] for r in records), default=None),
                   preview=records[:8], file_sha256=hashlib.sha256(content).hexdigest())
    bid = uuid.uuid4().hex
    with connection() as c:
        c.execute('INSERT INTO batches VALUES(?,?,?,?,?,?,?,?,?)',
                  (bid, did, Path(filename).name[:150], summary['file_sha256'], json.dumps(records),
                   json.dumps(summary), 'preview', now(), uid))
    return dict(id=bid, **summary)


def confirm(bid, uid):
    with connection() as c:
        c.execute('BEGIN IMMEDIATE')
        b = row(c, 'SELECT * FROM batches WHERE id=?', (bid,))
        if not b:
            raise ValueError('Importación no encontrada.')
        if b['status'] == 'confirmed':
            return dict(id=bid, reused=True)
        if dataset(c, b['dataset'])['archived']:
            raise ValueError('El historial está archivado. La importación no se confirmó.')
        summary = json.loads(b['summary'])
        if summary['invalid']:
            raise ValueError('Corrige las filas con errores antes de confirmar.')
        if not summary['accepted']:
            raise ValueError('El archivo no contiene observaciones nuevas para confirmar.')
        for r in json.loads(b['records']):
            old = row(c, 'SELECT units FROM observations WHERE dataset=? AND sku=? AND period=?',
                      (b['dataset'], r['sku'], r['period']))
            if old and old['units'] != r['units']:
                raise ValueError('El historial cambió desde la vista previa. Vuelve a importar el archivo.')
            c.execute('INSERT OR IGNORE INTO observations VALUES(?,?,?,?,?)',
                      (b['dataset'], r['sku'], r['period'], r['units'], bid))
        c.execute("UPDATE batches SET status='confirmed' WHERE id=?", (bid,))
        event(c, b['dataset'], 'importar', {'batch': bid, 'rows': summary['accepted']}, uid)
    return dict(id=bid)


def example_path(name):
    p = source()
    entry = next((f for f in (p or {}).get('files', []) if f['name'] == name), None)
    if not entry:
        raise ValueError('Archivo de demostración no encontrado. Ejecuta scripts/prepare_real_ml.py.')
    path = EXAMPLES/name
    if hashlib.sha256(path.read_bytes()).hexdigest() != entry['sha256']:
        raise ValueError('El archivo de demostración cambió respecto a su procedencia.')
    return path


def simulation_source():
    path = SIMULATION/'provenance.json'
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else None


def simulation_example_path(name):
    manifest = simulation_source()
    entry = next((f for f in (manifest or {}).get('stages', []) if f['name'] == name), None)
    if not entry:
        raise ValueError('Archivo de ensayo no encontrado. Ejecuta scripts/prepare_ml_simulation.py.')
    path = SIMULATION/name
    if hashlib.sha256(path.read_bytes()).hexdigest() != entry['sha256']:
        raise ValueError('El archivo de ensayo cambió respecto a su procedencia.')
    return path


def simulation(uid):
    manifest = simulation_source()
    if not manifest or manifest.get('kind') != 'synthetic_calibrated_trial':
        raise ValueError('Prepara el ensayo con scripts/prepare_ml_simulation.py.')
    path = simulation_example_path('01_historial_90_dias.csv')
    if manifest['end'] >= now()[:10]:
        raise ValueError('El ensayo contiene un día aún no cerrado. Vuelve a prepararlo.')
    source_text = ('ENSAYO SINTÉTICO calibrado con Car Parts (Zenodo DOI '
                   f"{manifest['source_doi']}). SKU y fechas asignados para prueba; "
                   'no son ventas de RODA.')
    did = create('RODA · ensayo diario calibrado', 'daily', source_text, uid)
    with connection() as c:
        c.execute('INSERT INTO scenarios VALUES(?,?,?)',
                  (did, 'roda-calibrated', hashlib.sha256((SIMULATION/'provenance.json').read_bytes()).hexdigest()))
    batch = preview(did, path.read_bytes(), path.name, None, uid)
    if batch['invalid'] or batch['accepted'] != manifest['stages'][0]['rows']:
        raise ValueError('El bloque inicial del ensayo no pasó la validación.')
    confirm(batch['id'], uid)
    return {'id': did}


def reset_simulation(did, uid):
    with connection() as c:
        d = dataset(c, did)
        scenario = row(c, 'SELECT kind FROM scenarios WHERE dataset=?', (did,))
        if d['archived'] or not scenario or scenario['kind'] != 'roda-calibrated':
            raise ValueError('Selecciona un ensayo activo para restablecerlo.')
        if c.execute("SELECT 1 FROM jobs WHERE dataset=? AND status='running'", (did,)).fetchone():
            raise ValueError('Espera a que finalice el cálculo.')
    result = simulation(uid)
    with connection() as c:
        c.execute('UPDATE datasets SET archived=1 WHERE id=?', (did,))
        event(c, did, 'archivar_ensayo', {'replacement': result['id']}, uid)
    return result


def demo(uid):
    if not source():
        raise ValueError('Prepara los datos con scripts/prepare_real_ml.py.')
    did = create('Car Parts · demostración mensual', 'monthly', source()['source_url'], uid, True)
    p = example_path('01_historial_inicial.csv')
    b = preview(did, p.read_bytes(), p.name, None, uid)
    confirm(b['id'], uid)
    return dict(id=did)


def reset(did, uid):
    with connection() as c:
        d = dataset(c, did)
        if d['archived']:
            raise ValueError('La demostración ya está archivada.')
        if not d['demo']:
            raise ValueError('Solo puede restablecerse la demostración de Car Parts.')
        if c.execute("SELECT 1 FROM jobs WHERE dataset=? AND status='running'", (did,)).fetchone():
            raise ValueError('Espera a que finalice el cálculo.')
    result = demo(uid)
    with connection() as c:
        c.execute('UPDATE datasets SET archived=1 WHERE id=?', (did,))
        event(c, did, 'archivar_demostracion', {'replacement': result['id']}, uid)
    return result


def listing(did=None):
    with connection() as c:
        datasets = [dict(r) for r in c.execute('SELECT * FROM datasets ORDER BY created_at DESC,id DESC')]
        readiness = {r['dataset']: dict(r) for r in c.execute('''
            SELECT v.dataset, COUNT(*) version_count,
                   (SELECT COUNT(*) FROM jobs j WHERE j.dataset=v.dataset AND j.action='compare' AND j.status='done') comparison_count
            FROM versions v GROUP BY v.dataset''')}
        for item in datasets:
            item['version_count'] = readiness.get(item['id'], {}).get('version_count', 0)
            item['comparison_count'] = readiness.get(item['id'], {}).get('comparison_count', 0)
        if not did:
            did = next((d['id'] for d in datasets if not d['archived']), None)
        if not did:
            return dict(datasets=datasets, source=source(), selected=None)
        d = dataset(c, did)
        scenario = row(c, 'SELECT kind,provenance_hash FROM scenarios WHERE dataset=?', (did,))
        d['scenario'] = scenario['kind'] if scenario else None
        stats = dict(c.execute('SELECT COUNT(*) rows,COUNT(DISTINCT sku) products,MIN(period) start,MAX(period) end,SUM(units IS NULL) missing FROM observations WHERE dataset=?', (did,)).fetchone())
        versions = [dict(r) for r in c.execute('SELECT * FROM versions WHERE dataset=? ORDER BY created_at,id', (did,))]
        for v in versions:
            v['report'] = json.loads(v['report'])
            del v['bundle']
        batches = [dict(r) for r in c.execute("SELECT id,filename,summary,status,created_at FROM batches WHERE dataset=? AND status='confirmed' ORDER BY created_at DESC", (did,))]
        for b in batches:
            b['summary'] = json.loads(b['summary'])
        jobs = [dict(r) for r in c.execute('SELECT * FROM jobs WHERE dataset=? ORDER BY created_at DESC,id DESC LIMIT 30', (did,))]
        for j in jobs:
            if j['result']:
                value = json.loads(j['result'])
                j['result'] = {k: v for k, v in value.items() if k != 'items'}
        events = [dict(r) for r in c.execute('SELECT * FROM activity WHERE dataset=? ORDER BY id DESC LIMIT 20', (did,))]
    return dict(datasets=datasets, selected=d, stats=stats, versions=versions, batches=batches, jobs=jobs, activity=events,
                source=source(), simulation=simulation_source() if scenario else None)


def start_job(did, action, uid, version_ids=None, managed=False):
    if did == 'roda-operacion-diaria' and not managed:
        raise ValueError('Usa «Revisar ahora» en la actualización automática para conservar los controles de calidad.')
    if action not in ['train', 'predict', 'compare']:
        raise ValueError('Acción no reconocida.')
    if not LOCK.acquire(blocking=False):
        raise ValueError('Hay un cálculo en curso. Espera a que finalice.')
    try:
        with connection() as c:
            d = dataset(c, did)
            if d['archived']:
                raise ValueError('Selecciona un historial activo.')
            records = [dict(r) for r in c.execute('SELECT sku,period,units FROM observations WHERE dataset=? ORDER BY sku,period', (did,))]
            if action == 'train' and d['demo'] and any(r['period'] > '2001-09' for r in records):
                raise ValueError('Las ventas de octubre de 2001 a marzo de 2002 están reservadas para comparar V1 y V2. No entrenes otra versión con ese bloque.')
            scenario = row(c, 'SELECT kind FROM scenarios WHERE dataset=?', (did,))
            if action == 'train' and scenario and scenario['kind'] == 'roda-calibrated':
                cutoff = simulation_source()['stages'][1]['end']
                if any(r['period'] > cutoff for r in records):
                    raise ValueError('Los últimos 30 días están reservados para evaluar V1 y V2. No entrenes con ese bloque.')
            if managed and action == 'train':
                required = {(date.fromisoformat(now()[:10])-timedelta(days=i)).isoformat() for i in range(1, 91)}
                observed = {}
                for r in records:
                    if r['units'] is not None:
                        observed.setdefault(r['sku'], set()).add(r['period'])
                eligible = {sku for sku, days in observed.items() if required <= days}
                records = [r for r in records if r['sku'] in eligible]
            if not records:
                raise ValueError('Primero importa el historial.')
            jid = uuid.uuid4().hex
            request = dict(action=action, frequency=d['frequency'], records=records)
            request['snapshot_rows'] = len(records)
            data_hash = hashlib.sha256(json.dumps(records, sort_keys=True).encode()).hexdigest()
            if action == 'train':
                if c.execute('SELECT 1 FROM versions WHERE dataset=? AND data_hash=?', (did, data_hash)).fetchone():
                    raise ValueError('Este historial ya tiene una versión entrenada. Importa periodos nuevos.')
                request['bundle'] = str(FOLDER/(jid+'.joblib'))
            else:
                ids = version_ids or ([d['active_version']] if d['active_version'] else [])
                expected = 1 if action != 'compare' or (managed and not d['active_version']) else 2
                if not isinstance(ids, list) or not all(isinstance(v, str) for v in ids) or len(ids) != expected or len(set(ids)) != len(ids):
                    raise ValueError('Selecciona una versión para predecir o las versiones necesarias para comparar.')
                specs = []
                for vid in ids:
                    v = row(c, 'SELECT * FROM versions WHERE id=? AND dataset=?', (vid, did))
                    if not v:
                        raise ValueError('La versión no pertenece a este historial.')
                    specs.append(dict(id=vid, bundle=v['bundle'], sha256=json.loads(v['report'])['model_sha256']))
                if action == 'compare':
                    request['versions'] = specs
                else:
                    request.update(specs[0])
            c.execute('INSERT INTO jobs VALUES(?,?,?,?,NULL,NULL,?,NULL,?)', (jid, did, action, 'running', now(), uid))
            if managed:
                c.execute('UPDATE automation SET pending_job=? WHERE id=1', (jid,))
            event(c, did, action, {'job': jid, 'data_hash': data_hash}, uid)
        threading.Thread(target=worker, args=(jid, did, request, data_hash), daemon=True).start()
        return dict(id=jid)
    except Exception:
        LOCK.release()
        raise


def worker(jid, did, request, data_hash):
    try:
        input_path, output_path = FOLDER/(jid+'.input.json'), FOLDER/(jid+'.output.json')
        input_path.write_text(json.dumps(request), encoding='utf-8')
        process = subprocess.run([str(runtime()), '-m', 'ml.lab_engine', str(input_path), str(output_path)],
            cwd=ROOT, capture_output=True, text=True, encoding='utf-8', timeout=900,
            env=dict(os.environ, PYTHONUTF8='1', LOKY_MAX_CPU_COUNT='2'),
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        result = json.loads(output_path.read_text(encoding='utf-8')) if output_path.exists() else {}
        if process.returncode or 'error' in result:
            raise ValueError(result.get('error', 'No se pudo ejecutar el modelo. Revisa el entorno de ML.'))
        result['data_sha256'] = data_hash
        result['snapshot_rows'] = request['snapshot_rows']
        with connection() as c:
            if request['action'] == 'train':
                count = c.execute('SELECT COUNT(*) FROM versions WHERE dataset=?', (did,)).fetchone()[0]
                c.execute('INSERT INTO versions VALUES(?,?,?,?,?,?,?)',
                          (jid, did, 'V'+str(count+1), request['bundle'], json.dumps(result), data_hash, now()))
            c.execute("UPDATE jobs SET status='done',result=?,finished_at=? WHERE id=?", (json.dumps(result), now(), jid))
    except Exception as exc:
        with connection() as c:
            c.execute("UPDATE jobs SET status='error',error=?,finished_at=? WHERE id=?", (str(exc)[:1000], now(), jid))
    finally:
        LOCK.release()


def activate(did, vid, uid, managed=False):
    if did == 'roda-operacion-diaria' and not managed:
        raise ValueError('La versión operativa se activa mediante la revisión automática y sus controles de calidad.')
    with connection() as c:
        d = dataset(c, did)
        if d['archived']:
            raise ValueError('El historial está archivado.')
        if not row(c, 'SELECT id FROM versions WHERE dataset=? AND id=?', (did, vid)):
            raise ValueError('Versión no encontrada.')
        if vid == d['active_version']:
            return dict(id=vid)
        version = row(c, 'SELECT report FROM versions WHERE id=? AND dataset=?', (vid, did))
        report = json.loads(version['report'])
        comparison = row(c, "SELECT result FROM jobs WHERE dataset=? AND action='compare' AND status='done' ORDER BY created_at DESC,id DESC LIMIT 1", (did,))
        if not comparison:
            raise ValueError('Primero compara las versiones con ventas posteriores antes de activar una.')
        observed = json.loads(comparison['result'])
        current_rows = c.execute('SELECT COUNT(*) FROM observations WHERE dataset=?', (did,)).fetchone()[0]
        if observed.get('snapshot_rows') != current_rows:
            raise ValueError('Hay ventas nuevas desde la comparación. Vuelve a evaluar las versiones.')
        scores = {s['version']: s['mae'] for s in observed['scores']}
        if vid not in scores or scores[vid] >= observed['baseline']['mae'] * .99:
            raise ValueError('La versión no supera al promedio histórico por al menos 1% en ventas posteriores.')
        if vid != observed.get('recommended'):
            raise ValueError('Activa únicamente la versión recomendada por la comparación actual.')
        if not d['active_version']:
            if (not report['candidates'] or min(s['mae'] for s in report['candidates'])
                    >= report['validation_baseline']['mae'] * .99):
                raise ValueError('La versión no reduce al menos 1% el error frente al promedio histórico en validación.')
        else:
            if (vid not in scores or d['active_version'] not in scores
                    or scores[vid] >= scores[d['active_version']] * .99
                    or scores[vid] >= observed['baseline']['mae'] * .99):
                raise ValueError('La candidata debe reducir al menos 1% el MAE frente a la activa y al promedio histórico.')
        c.execute('UPDATE datasets SET active_version=? WHERE id=?', (vid, did))
        event(c, did, 'activar_version', {'previous': d['active_version'], 'version': vid}, uid)
    return dict(id=vid)


def results(jid, q='', page=1):
    with connection() as c:
        j = row(c, 'SELECT * FROM jobs WHERE id=?', (jid,))
    if not j or not j['result']:
        raise ValueError('Resultado todavía no disponible.')
    result = json.loads(j['result'])
    if 'items' in result:
        items = [r for r in result['items'] if q.lower() in r['sku'].lower()]
        result.update(items=items[(max(1,page)-1)*20:max(1,page)*20], total=len(items))
    return result


def version_file(vid):
    with connection() as c:
        version = row(c, 'SELECT * FROM versions WHERE id=?', (vid,))
    if not version:
        raise ValueError('Versión no encontrada.')
    path = Path(version['bundle']).resolve()
    if not path.is_relative_to(FOLDER.resolve()) or not path.is_file():
        raise ValueError('El archivo registrado de la versión no está disponible.')
    report = json.loads(version['report'])
    if hashlib.sha256(path.read_bytes()).hexdigest() != report['model_sha256']:
        raise ValueError('El modelo cambió respecto al hash registrado; no se puede descargar.')
    return path, version['label']


def export(jid):
    with connection() as c:
        j = row(c, 'SELECT result FROM jobs WHERE id=?', (jid,))
    result = json.loads(j['result'] or '{}') if j else {}
    if 'items' not in result:
        raise ValueError('Este cálculo no contiene pronósticos para exportar.')
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=['sku', 'period', 'prediction', 'baseline', 'missing', 'version', 'cutoff', 'frequency', 'data_sha256', 'model_sha256'])
    writer.writeheader()
    writer.writerows(dict(r, **{k: result.get(k, '') for k in ['version','cutoff','frequency','data_sha256','model_sha256']}) for r in result['items'])
    return '\ufeff'+out.getvalue()
